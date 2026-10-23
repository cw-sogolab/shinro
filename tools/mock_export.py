#!/usr/bin/env python3
"""모의고사 분석 데이터 만들기 (진학 상담 「모의고사 분석」 탭).

입력: iCloud 일공/자료/EJU성적/모의고사 성적/
      2026_{N}회_전국모의고사_채점시트.xlsx  (검수 끝난 채점시트)
      EJU_점수예측_분석.xlsx                  (난이도분석 시트)
출력: docs/assets/mock.js  (window.MOCK)

학생 이름·수험번호·아이디는 출력에 넣지 않는다. 집계(인원·평균·분포·문항 정답률·지망대학 건수)만.
같은 학생의 회차 간 변화는 102classID로 짝을 지은 뒤 집계만 남긴다.

사용:
  python3 tools/mock_export.py --dir "<모의고사 성적 폴더>" --year 2026 --rounds 1 2 3 4
  오래 걸리면 회차별로: --stage round --rounds 1  (… 4까지)  →  --stage final
"""
import argparse
import hashlib
import json
import os
import statistics
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "assets" / "mock.js"

ITEM_SHEETS = {"jr": "독해", "jl": "청해청독해", "so": "종합과목", "phy": "물리", "chem": "화학", "bio": "생물"}
ITEM_LABEL = {"jr": "독해", "jl": "청해·청독해", "so": "종합과목", "phy": "물리", "chem": "화학", "bio": "생물"}
SCORE_COLS = {"jr": ["독해"], "jl": ["청독해", "청해청독해"], "ki": ["기술"], "so": ["종과", "종합과목"],
              "phy": ["물리"], "chem": ["화학"], "bio": ["생물"], "m1": ["코스1"], "m2": ["코스2"]}
MAX = {"jp": 400, "jr": 200, "jl": 200, "ki": 50, "so": 200, "m1": 200, "m2": 200, "phy": 100, "chem": 100, "bio": 100,
       "bunka": 800, "rika": 800}
STEP = {"jp": 20, "jr": 10, "jl": 10, "ki": 5, "so": 10, "m1": 10, "m2": 10, "phy": 5, "chem": 5, "bio": 5, "bunka": 40, "rika": 40}


def nfc(s):
    return unicodedata.normalize("NFC", str(s))


def resolve(path):
    """iCloud 한글 경로(NFD)를 NFC로 찾아 실제 경로를 돌려준다."""
    p = Path(os.path.expanduser(path))
    if p.exists():
        return p
    cur = Path(p.anchor)
    for part in p.parts[1:]:
        hit = [x for x in os.listdir(cur) if nfc(x) == nfc(part)]
        if not hit:
            raise FileNotFoundError(path)
        cur = cur / hit[0]
    return cur


def num(v):
    if v is None or v == "" or isinstance(v, bool):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f


def rows_of(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def find_sheet(wb, name):
    for s in wb.sheetnames:
        if nfc(s) == nfc(name):
            return wb[s]
    return None


def col_index(header, names):
    h = [nfc(x).replace("\n", "").strip() if x is not None else "" for x in header]
    for n in names:
        if n in h:
            return h.index(n)
    return None


def hid(x, salt):
    return hashlib.sha1((salt + str(x)).encode()).hexdigest()[:12]


def stats(vals, key):
    vals = [v for v in vals if v is not None]
    if not vals:
        return None
    vals.sort()
    mx, st = MAX[key], STEP[key]
    nb = mx // st
    hist = [0] * nb
    for v in vals:
        hist[min(nb - 1, max(0, int(v // st)))] += 1
    q = lambda p: vals[min(len(vals) - 1, int(round(p * (len(vals) - 1))))]
    return {"n": len(vals), "mean": round(statistics.mean(vals), 1), "sd": round(statistics.pstdev(vals), 1),
            "q": [q(.1), q(.25), q(.5), q(.75), q(.9)], "hist": hist, "step": st, "max": mx,
            "sorted": [round(v) for v in vals]}


# ---------------- 회차 하나 ----------------
def do_round(dirp, year, n, work, salt):
    f = resolve(f"{dirp}/{year}_{n}회_전국모의고사_채점시트.xlsx")
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    warn = []

    # 1) 종합점수: 학생별 점수·캠퍼스·지망
    ws = find_sheet(wb, "종합점수(지망대학포함)")
    R = rows_of(ws)
    head = R[0]
    ci = {k: col_index(head, v) for k, v in SCORE_COLS.items()}
    c_camp = col_index(head, ["캠퍼스구분", "캠퍼스"])
    c_kind = col_index(head, ["문/이과"])
    c_issue = col_index(head, ["시험장 ISSUE"])
    c_w = [col_index(head, [f"지망대학{i}"]) for i in (1, 2, 3)]
    c_wf = [col_index(head, [f"지망학부{i}", f"지망학과{i}"]) for i in (1, 2, 3)]
    c_r = [col_index(head, [f"추천학교{i}"]) for i in (1, 2, 3)]
    for k, v in ci.items():
        if v is None:
            warn.append(f"{n}회 종합점수: '{SCORE_COLS[k][0]}' 열 없음")

    # 수험번호 → 102classID (회차 간 짝짓기용, 출력하지 않음)
    idmap = {}
    wsi = find_sheet(wb, f"학생정보_{year}_{n}회")
    if wsi is not None:
        IR = rows_of(wsi)
        a, b = col_index(IR[0], ["수험번호"]), col_index(IR[0], ["102classID"])
        if a is not None and b is not None:
            for r in IR[1:]:
                if r[a] is not None and r[b]:
                    idmap[str(int(num(r[a])) if num(r[a]) is not None else r[a])] = str(r[b]).strip().lower()
    else:
        warn.append(f"{n}회: 학생정보_{year}_{n}회 시트 없음 — 성장 추이 짝짓기 불가")

    students = []
    for r in R[1:]:
        sid = r[0]
        if sid is None or num(sid) is None:
            continue
        sc = {k: (num(r[c]) if c is not None else None) for k, c in ci.items()}
        if all(v is None for v in sc.values()):
            continue
        for k in sc:
            if sc[k] is not None and not (0 <= sc[k] <= MAX[k]):
                sc[k] = None  # 범위 밖(입력 오류)은 뺀다
        jp = sc["jr"] + sc["jl"] if sc["jr"] is not None and sc["jl"] is not None else None
        sci = [sc[k] for k in ("phy", "chem", "bio") if sc[k] is not None]
        kind = nfc(r[c_kind]).strip() if c_kind is not None and r[c_kind] else ""
        if kind not in ("문과", "이과"):
            kind = "이과" if len(sci) >= 1 or sc["m2"] is not None else ("문과" if sc["so"] is not None or sc["m1"] is not None else "일본어만")
        bunka = jp + sc["so"] + sc["m1"] if kind == "문과" and None not in (jp, sc["so"], sc["m1"]) else None
        rika = jp + sum(sorted(sci)[-2:]) + sc["m2"] if kind == "이과" and jp is not None and len(sci) >= 2 and sc["m2"] is not None else None
        camp = nfc(r[c_camp]).strip() if c_camp is not None and r[c_camp] else "미기재"
        wishes = [(nfc(r[c]).strip() if c is not None and r[c] else None) for c in c_w]
        wfac = [(nfc(r[c]).strip() if c is not None and r[c] else None) for c in c_wf]
        recs = [(nfc(r[c]).strip() if c is not None and r[c] else None) for c in c_r]
        key = str(int(num(sid)))
        cid = idmap.get(key)
        students.append({"h": hid(cid, salt) if cid else None, "camp": camp, "kind": kind, "jp": jp, "bunka": bunka, "rika": rika,
                         **sc, "w": wishes, "wf": wfac, "rec": recs})

    # 2) 문항 분석
    items = {}
    for key, sname in ITEM_SHEETS.items():
        ws = find_sheet(wb, sname)
        if ws is None:
            warn.append(f"{n}회: '{sname}' 시트 없음")
            continue
        IR = rows_of(ws)
        if len(IR) < 3:
            continue
        h2 = IR[1]
        # 응답 블록: 5번째 칸부터 숫자 머리글이 이어지는 동안
        N = 0
        for v in h2[4:]:
            if num(v) is None:
                break
            N += 1
        if not N:
            warn.append(f"{n}회 {sname}: 문항 머리글을 못 찾음")
            continue
        rs, cs = 4, 4 + N + 1
        keyrow = [num(v) for v in IR[0][cs:cs + N]]
        recs = []
        for r in IR[2:]:
            sc = num(r[1]) if len(r) > 1 else None
            resp = [num(v) for v in r[rs:rs + N]]
            corr = [num(v) for v in r[cs:cs + N]]
            if sc is None or all(v is None for v in resp):
                continue
            recs.append((sc, resp, corr))
        if not recs:
            continue
        recs.sort(key=lambda x: -x[0])
        g = max(1, round(len(recs) * 0.27))
        up, lo = recs[:g], recs[-g:]
        out = []
        for i in range(N):
            ok = [c[2][i] for c in recs if c[2][i] is not None]
            p = sum(ok) / len(ok) if ok else None
            pu = sum(c[2][i] or 0 for c in up) / len(up)
            pl = sum(c[2][i] or 0 for c in lo) / len(lo)
            opts = {}
            for c in recs:
                v = c[1][i]
                k2 = "무응답" if v is None else str(int(v))
                opts[k2] = opts.get(k2, 0) + 1
            kk = keyrow[i]
            wrong = sorted(((k2, c) for k2, c in opts.items() if k2 != "무응답" and kk is not None and k2 != str(int(kk))), key=lambda x: -x[1])
            out.append({"q": i + 1, "key": int(kk) if kk is not None else None, "p": round(p, 3) if p is not None else None,
                        "d": round(pu - pl, 3), "opts": opts, "lure": wrong[0][0] if wrong else None,
                        "lureP": round(wrong[0][1] / len(recs), 3) if wrong else None})
        items[key] = {"label": ITEM_LABEL[key], "n": len(recs), "items": out}

    res = {"n": n, "students": students, "items": items, "warn": warn}
    (work / f"round{n}.json").write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
    print(f"{n}회: 학생 {len(students)}명, 문항 시트 {len(items)}개, 경고 {len(warn)}")
    for w in warn:
        print("  ", w)


# ---------------- 합치기 ----------------
def difficulty(dirp):
    f = resolve(f"{dirp}/EJU_점수예측_분석.xlsx")
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = find_sheet(wb, "난이도분석")
    R = rows_of(ws)
    hi = next(i for i, r in enumerate(R) if r and nfc(r[0] or "") == "연도")
    head = [nfc(x or "") for x in R[hi]]
    out = []
    for r in R[hi + 1:]:
        if not r or r[0] is None or num(r[0]) is None:
            continue
        d = dict(zip(head, r))
        out.append({"y": int(num(d.get("연도"))), "r": nfc(d.get("회차") or ""), "k": nfc(d.get("계열") or ""), "s": nfc(d.get("과목") or ""),
                    "vs": nfc(d.get("비교 EJU") or ""), "n": int(num(d.get("짝 인원")) or 0),
                    "mock": round(num(d.get("모의 평균")) or 0, 1), "eju": round(num(d.get("EJU 평균")) or 0, 1),
                    "diff": round(num(d.get("차이(모의−EJU)")) or 0, 1), "rr": round(num(d.get("상관 r")) or 0, 2),
                    "j": nfc(d.get("판정") or ""), "delta": None if num(d.get("2026−2025 난이도차 δ")) is None else round(num(d.get("2026−2025 난이도차 δ")), 1),
                    "dj": nfc(d.get("δ 해석") or "")})
    return out


def final(dirp, year, rounds, work):
    RS = [json.loads((work / f"round{n}.json").read_text(encoding="utf-8")) for n in rounds]
    out_rounds = []
    for R in RS:
        st = R["students"]
        camps = {}
        for s in st:
            camps[s["camp"]] = camps.get(s["camp"], 0) + 1
        kinds = {}
        for s in st:
            kinds[s["kind"]] = kinds.get(s["kind"], 0) + 1
        subj = {}
        for k in ("jp", "jr", "jl", "ki", "so", "m1", "m2", "phy", "chem", "bio", "bunka", "rika"):
            v = stats([s.get(k) for s in st], k)
            if v:
                subj[k] = v
        bycamp = {}
        for c in camps:
            ss = [s for s in st if s["camp"] == c]
            bycamp[c] = {k: (round(statistics.mean(x), 1) if x else None) for k in ("jp", "so", "m1", "bunka", "rika")
                         for x in [[s[k] for s in ss if s.get(k) is not None]]}
            bycamp[c]["n"] = len(ss)
        # 지망대학
        w1, wall, wfac = {}, {}, {}
        hit = tot = 0
        for s in st:
            ws = [w for w in s["w"] if w]
            if s["w"][0]:
                w1[s["w"][0]] = w1.get(s["w"][0], 0) + 1
                f = (s["wf"][0] or "").strip()
                if f:
                    kf = s["w"][0] + " " + f
                    wfac[kf] = wfac.get(kf, 0) + 1
            for w in ws:
                wall[w] = wall.get(w, 0) + 1
            recs = [x for x in s["rec"] if x]
            if s["w"][0] and recs:
                tot += 1
                hit += s["w"][0] in recs
        top = lambda d, k: [[a, b] for a, b in sorted(d.items(), key=lambda x: -x[1])[:k]]
        out_rounds.append({"n": R["n"], "takers": len(st), "camps": camps, "kinds": kinds, "subj": subj, "bycamp": bycamp,
                           "items": R["items"], "wish1": top(w1, 25), "wishAll": top(wall, 25), "wishFac": top(wfac, 20),
                           "wishInRec": {"n": tot, "hit": hit}, "warn": R["warn"]})
    # 성장 추이: 102classID로 짝지은 학생만, 집계만 남김
    by = {}
    for R in RS:
        for s in R["students"]:
            if s["h"]:
                by.setdefault(s["h"], {})[R["n"]] = s
    growth = {}
    for k in ("jp", "so", "m1", "bunka"):
        both = [(v[rounds[0]][k], v[rounds[-1]][k]) for v in by.values() if rounds[0] in v and rounds[-1] in v
                and v[rounds[0]].get(k) is not None and v[rounds[-1]].get(k) is not None]
        allr = [v for v in by.values() if all(n in v and v[n].get(k) is not None for n in rounds)]
        if not both:
            continue
        ch = [b - a for a, b in both]
        step = 10 if k != "bunka" else 20
        lo = int(min(ch) // step * step)
        hist = {}
        for c in ch:
            b = int(c // step * step)
            hist[b] = hist.get(b, 0) + 1
        cohort = {}
        for camp in sorted({v[rounds[-1]]["camp"] for v in allr}):
            vv = [v for v in allr if v[rounds[-1]]["camp"] == camp]
            if len(vv) >= 5:  # 5명 미만 집단은 내지 않음
                cohort[camp] = {"n": len(vv), "mean": [round(statistics.mean(v[n][k] for v in vv), 1) for n in rounds]}
        growth[k] = {"pairs": len(both), "meanChange": round(statistics.mean(ch), 1), "up": sum(c > 0 for c in ch),
                     "down": sum(c < 0 for c in ch), "step": step, "hist": sorted([[a, b] for a, b in hist.items()]),
                     "all": {"n": len(allr), "mean": [round(statistics.mean(v[n][k] for v in allr), 1) for n in rounds] if allr else []},
                     "cohort": cohort}
    payload = {"built": date.today().isoformat(), "year": year, "rounds": out_rounds, "growth": growth,
               "growthRounds": rounds, "difficulty": difficulty(dirp)}
    OUT.write_text("/* 자동 생성 — tools/mock_export.py. 이름·수험번호 없음(집계만). 그래도 학원 밖 공개 금지 */\n"
                   "window.MOCK = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(f"완료: {len(out_rounds)}개 회차 → {OUT.relative_to(ROOT)} ({OUT.stat().st_size // 1024} KB)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default="~/Library/Mobile Documents/com~apple~CloudDocs/일공/자료/EJU성적/모의고사 성적")
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument("--rounds", type=int, nargs="+", default=[1, 2, 3, 4])
    ap.add_argument("--stage", choices=["all", "round", "final"], default="all")
    ap.add_argument("--work", default="/tmp/mock_work")
    a = ap.parse_args()
    work = Path(a.work)
    work.mkdir(parents=True, exist_ok=True)
    salt = os.urandom(8).hex() if a.stage == "all" else (work / "salt").read_text() if (work / "salt").exists() else None
    if salt is None:
        salt = os.urandom(8).hex()
        (work / "salt").write_text(salt)
    if a.stage in ("all", "round"):
        for n in a.rounds:
            do_round(a.dir, a.year, n, work, salt)
    if a.stage in ("all", "final"):
        final(a.dir, a.year, a.rounds, work)


if __name__ == "__main__":
    main()
