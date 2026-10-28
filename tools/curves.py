#!/usr/bin/env python3
"""합격 곡선 만들기 (학교·학부·계열별 로지스틱 회귀)

입력 : EJU통합_성적분석_vN.xlsx 의 「데이터」 시트 (연도별 지원자 점수 + 합격여부)
출력 : ① 엑셀 EJU_합격곡선_vN.xlsx (곡선 계수·검증·계산기)
       ② docs/assets/curves.js (곡선 보기 화면용, 학생 개인 정보 없음 — 점수 구간별 인원만)

사용 : python3 tools/curves.py --src <통합본.xlsx> --xlsx <출력.xlsx> --js docs/assets/curves.js

방법 요약
 1. 지원자마다 응시 과목으로 계열을 정함: 문과+수학 / 문과 / 이과+수학 / 이과 / 일본어만. 총점에서 기술은 뺌.
 2. 학교·학부·계열마다 6월·11월·둘 중 높은 점수 중 합격을 가장 잘 가르는 회차(AUC)를 고름.
    합격·불합격이 각각 10명 이상일 때만 곡선을 만듦.
 3. 같은 점수의 합격률이 해마다 떨어지는 추세가 있어서, 최근 연도에 가중치를 주고(1·.7·.5·.35·.25)
    연도별 공통 보정값 δ를 함께 추정한 뒤, 다음 해 예측에는 가장 최근 해의 δ를 씀.
 4. 검증: (2022~2024 → 2025), (2022~2025 → 2026) 두 번 「미래를 모르는 상태」로 맞춰 봄.
 5. (v2) 데이터 정리: 이 데이터는 여러 학원이 기관에 제출한 것이라 믿기 어려운 기록이 섞여 있다. 곡선을 만들기 전에
    ① 같은 줄이 통째로 중복된 기록 제거, ② 영어 증빙(TOEFL·TOEIC)이 합격자 대부분에 있는 「영어 필수 학부」인데 증빙이 없는 기록 제거,
    ③ 학부·합격여부별로 점수가 너무 동떨어진 값(중앙값에서 3.5 표준편차 상당 이상)만 그 점수에서 제외. 제거 건수는 CLEAN 에 기록.
"""
import argparse, os, unicodedata, datetime
import numpy as np, pandas as pd
import yoko_rule as yr   # 2027 모집요강이 인정하는 계열·과목 조합 (tools/yoko_rule.py)

SCI = ['물리', '화학', '생물']
MIN_EACH = 10
WEIGHTS = [1, .7, .5, .35, .25]          # 최근 연도부터
ZONES = [(0, .15, '15% 미만'), (.15, .4, '도전'), (.4, .7, '적정'), (.7, 1.01, '안정')]
sig = lambda t: 1 / (1 + np.exp(-t))
logit = lambda p: np.log(p / (1 - p))


def resolve(p):
    """macOS 한글 파일명(NFD)도 찾을 수 있게 경로를 한 단계씩 정규화해서 찾음"""
    p = os.path.expanduser(p)
    if os.path.exists(p): return p
    cur = '/' if p.startswith('/') else '.'
    for part in [x for x in p.split('/') if x]:
        hit = [x for x in os.listdir(cur) if unicodedata.normalize('NFC', x) == unicodedata.normalize('NFC', part)]
        if not hit: raise FileNotFoundError(p)
        cur = os.path.join(cur, hit[0])
    return cur


# ---------- 1. 준비 ----------
CLEAN = {}
ENG_SHARE, ENG_MIN = 0.6, 8     # 영어 필수 학부 판정: 합격자 중 TOEFL/TOEIC 기재 비율 ≥ 60% 이고 기재 합격자 ≥ 8명
OUT_Z, OUT_MIN = 3.5, 8         # 이상치: 로버스트 z(중앙값·MAD) ≥ 3.5, 그룹 크기 ≥ 8


def clean_rows(df):
    """믿기 어려운 기록 제거 (중복·영어 증빙 없음). 점수 이상치는 add_scores 뒤에 따로."""
    n0 = len(df); CLEAN['처음'] = n0
    df = df.drop_duplicates(subset=[c for c in df.columns if c not in ('y',)]).copy()
    CLEAN['중복'] = n0 - len(df)
    df['_eng'] = df['TOEFL_iBT'].notna() | df['TOEIC'].notna()
    g = df.groupby(['학교n', '학부n'])
    ok = g.apply(lambda x: bool(x['y'].sum() and x.loc[x.y == 1, '_eng'].mean() >= ENG_SHARE and x.loc[x.y == 1, '_eng'].sum() >= ENG_MIN), include_groups=False)
    req = ok[ok].index
    m = pd.MultiIndex.from_frame(df[['학교n', '학부n']]).isin(req) & ~df['_eng'].values
    CLEAN['영어필수_학부'] = int(len(req)); CLEAN['영어증빙없음'] = int(m.sum()); CLEAN['영어증빙없음_합격'] = int((m & (df['y'].values == 1)).sum())
    df = df[~m].drop(columns='_eng').copy()
    CLEAN['남음'] = len(df)
    return df


def drop_score_outliers(df):
    """학부·합격여부별로 점수가 너무 동떨어진 값은 그 점수만 비운다 (기록 전체를 지우지 않음)."""
    cols = [c for c in df.columns if c.startswith('S_')]; tot = 0
    g = df.groupby(['학교n', '학부n', 'y'])
    for c in cols:
        med = g[c].transform('median'); mad = g[c].transform(lambda s: (s - s.median()).abs().median()); cnt = g[c].transform('count')
        z = (df[c] - med).abs() / (1.4826 * mad.replace(0, np.nan))
        bad = (z >= OUT_Z) & (cnt >= OUT_MIN)
        tot += int(bad.sum()); df.loc[bad, c] = np.nan
    CLEAN['점수이상치(점수 칸)'] = tot
    return df


def load(src):
    df = pd.read_excel(resolve(src), sheet_name='데이터', engine='openpyxl', usecols=list(range(30)))
    last = int(df['연도'].max())
    df = df[(df['연도'] >= last - 4) & df['합격여부'].isin(['합격', '불합격'])].copy()
    df['y'] = (df['합격여부'] == '합격').astype(int)
    nk = lambda s: unicodedata.normalize('NFKC', str(s)) if pd.notna(s) else ''
    df['학교n'] = df['학교'].map(nk); df['학부n'] = df['학부'].map(nk)
    df = clean_rows(df)
    for ses in ['6월', '11월']:
        res = df.apply(lambda r: components(r, ses), axis=1)
        for j, c in enumerate(['kind', 'jp', 'so', 'm1', 'r', 'm2']):
            df[f'{c}_{ses}'] = [t[j] for t in res]
    return drop_score_outliers(add_scores(df))


def components(row, ses):
    """회차별: 계열(문과/이과/일본어만)과 과목별 점수"""
    g = lambda c: row.get(f'{ses}_{c}')
    jp, so, m1, m2 = g('일본어'), g('종합'), g('수학1'), g('수학2')
    sci = [g(c) for c in SCI if pd.notna(g(c))]
    if pd.isna(jp): kind = None
    elif pd.notna(so) and not sci: kind = '문과'
    elif len(sci) >= 2 and pd.isna(so): kind = '이과'
    elif pd.isna(so) and not sci and pd.isna(m1) and pd.isna(m2): kind = '일본어만'
    else: kind = None
    return kind, jp, so, m1, (sum(sci[:2]) if len(sci) >= 2 else np.nan), m2


# 과목 조합: 학부가 실제로 무엇을 보는지 모르므로 후보를 모두 시험해 합격을 가장 잘 가르는 것을 고름
SETS = {'문과': ['J', 'JS', 'JSM'], '이과': ['J', 'JR', 'JRM'], '일본어만': ['J']}
SET_PARTS = {'J': ['jp'], 'JS': ['jp', 'so'], 'JSM': ['jp', 'so', 'm1'], 'JR': ['jp', 'r'], 'JRM': ['jp', 'r', 'm2']}
SET_MAX = {'J': 400, 'JS': 600, 'JSM': 800, 'JR': 600, 'JRM': 800}
SET_LABEL = {'J': '일본어', 'JS': '일본어+종합', 'JSM': '일본어+종합+수학1', 'JR': '일본어+이과2과목', 'JRM': '일본어+이과2과목+수학2'}


def add_scores(df):
    """계열·과목조합·회차별 점수 열을 미리 만들어 둠 (S_문과_JSM_최고 등)"""
    for kind, sets in SETS.items():
        for sset in sets:
            one = {se: sum(df[f'{c}_{se}'] for c in SET_PARTS[sset]).where(df[f'kind_{se}'] == kind) for se in ['6월', '11월']}
            df[f'S_{kind}_{sset}_6월'] = one['6월']; df[f'S_{kind}_{sset}_11월'] = one['11월']
            df[f'S_{kind}_{sset}_최고'] = pd.concat([one['6월'], one['11월']], axis=1).max(axis=1)
    return df


def score_for(d, kind, sset, ses):
    return d[f'S_{kind}_{sset}_{ses}']


def auc(s, y):
    s = np.asarray(s, float); y = np.asarray(y)
    n1 = (y == 1).sum(); n0 = (y == 0).sum()
    if not n1 or not n0: return np.nan
    r = pd.Series(s).rank().values
    return (r[y == 1].sum() - n1 * (n1 + 1) / 2) / (n1 * n0)


def fit_logit(x, y, w=None, off=None, lam=1.0):
    """가중 로지스틱 회귀(뉴턴법). off = 연도 보정값. 기울기에 약한 규제(lam)."""
    w = np.ones_like(x) if w is None else w; off = np.zeros_like(x) if off is None else off
    mu, sd = x.mean(), x.std() or 1.0; z = (x - mu) / sd; c = np.zeros(2)
    for _ in range(60):
        p = sig(c[0] + c[1] * z + off); r = w * (p - y)
        g = np.array([r.sum(), (r * z).sum() + lam * c[1]]); W = w * p * (1 - p)
        H = np.array([[W.sum(), (W * z).sum()], [(W * z).sum(), (W * z * z).sum() + lam]])
        st = np.linalg.solve(H, g); c -= st
        if np.abs(st).max() < 1e-9: break
    b = c[1] / sd
    return c[0] - b * mu, b


# ---------- 2. 곡선 ----------
def build_units(train):
    units = []
    for (sch, fac), d in train.groupby(['학교n', '학부n']):
        rule = yr.rule(sch, fac)   # 요강이 있는 학부는 요강 계열·조합만 (없으면 예전 방식)
        for kind, sets in SETS.items():
            if rule:
                if kind not in rule['kinds']:
                    continue
                sets = [x for x in sets if x in rule['sets'][kind]]
            cand = []
            for sset in sets:
                for ses in ['최고', '6월', '11월']:
                    s = score_for(d, kind, sset, ses); m = s.notna()
                    n1 = int(d.loc[m, 'y'].sum()); n0 = int(m.sum() - n1)
                    if n1 >= MIN_EACH and n0 >= MIN_EACH:
                        cand.append((auc(s[m], d.loc[m, 'y']), sset, ses, int(m.sum()), n1))
            if not cand: continue
            # AUC가 같으면(소수 둘째 자리) 과목이 많은 쪽 → 최고 회차 우선
            cand.sort(key=lambda c: (-round(c[0], 2), -len(SET_PARTS[c[1]]), ['최고', '6월', '11월'].index(c[2])))
            A, sset, ses, n, n1 = cand[0]
            units.append(dict(학교=sch, 학부=fac, 계열=kind, 요강=bool(rule), 과목=sset, 회차=ses, n=n, 합격=n1, AUC=A, a=0.0, b=0.0))
    return pd.DataFrame(units)


def unit_rows(d, U):
    out = []; g = d.groupby(['학교n', '학부n'])
    for i, u in U.iterrows():
        try: dd = g.get_group((u['학교'], u['학부']))
        except KeyError: continue
        s = score_for(dd, u['계열'], u['과목'], u['회차']); m = s.notna()
        out.append(pd.DataFrame(dict(u=i, s=s[m].values.astype(float), y=dd.loc[m, 'y'].values.astype(float), yr=dd.loc[m, '연도'].values)))
    return pd.concat(out, ignore_index=True) if out else pd.DataFrame(columns=['u', 's', 'y', 'yr'])


def year_delta(R, U, years):
    eta = U.loc[R.u, 'a'].values + U.loc[R.u, 'b'].values * R.s.values; D = {}
    for yv in years:
        m = R.yr.values == yv; e = eta[m]; tgt = R.y.values[m].sum(); dl = 0.0
        for _ in range(60):
            p = sig(e + dl); dl -= (p.sum() - tgt) / max((p * (1 - p)).sum(), 1e-9)
        D[yv] = dl
    return D


def fit_all(df, years, adjust=True):
    """years 로 학습. 반환: 곡선표 U, 다음 해에 쓸 보정값 shift, 연도별 δ"""
    tr = df[df['연도'].isin(years)]
    U = build_units(tr); R = unit_rows(tr, U)
    last = max(years); wts = {last - k: WEIGHTS[k] for k in range(len(WEIGHTS))}
    D = {y: 0.0 for y in years}
    for _ in range(3 if adjust else 1):
        for i in U.index:
            m = R.u.values == i
            w = np.array([wts[v] for v in R.yr.values[m]]) if adjust else None
            off = np.array([D[v] for v in R.yr.values[m]]) if adjust else None
            U.loc[i, ['a', 'b']] = fit_logit(R.s.values[m], R.y.values[m], w, off)
        if not adjust: break
        D2 = year_delta(R, U, years); D = {k: D[k] + v for k, v in D2.items()}
    U = U[U.b > 0].reset_index(drop=True)
    return U, (D[last] if adjust else 0.0), D


def predict(df, U, shift, year):
    T = unit_rows(df[df['연도'] == year], U)
    T['p'] = sig(U.loc[T.u, 'a'].values + U.loc[T.u, 'b'].values * T.s.values + shift)
    return T


def zone_table(T):
    rows = []
    for lo, hi, name in ZONES:
        s = T[(T.p >= lo) & (T.p < hi)]
        rows.append(dict(구간=name, 범위=f'{lo:.0%}~{min(hi, 1):.0%}', 인원=len(s), 예측평균=s.p.mean(), 실제합격률=s.y.mean()))
    return pd.DataFrame(rows)


# ---------- 3. 실행 ----------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', required=True); ap.add_argument('--xlsx', required=True); ap.add_argument('--js', required=True)
    ap.add_argument('--version', default='v1')
    ap.add_argument('--stage', default='all', choices=['all', 'prep', 'val', 'final'], help='오래 걸리면 단계별로 나눠 실행')
    ap.add_argument('--work', default='/tmp/curves_work')
    A = ap.parse_args(); os.makedirs(A.work, exist_ok=True); import pickle
    pk = lambda n: os.path.join(A.work, n + '.pkl')
    if A.stage in ('all', 'prep'):
        df = load(A.src); df.to_pickle(pk('prep')); pickle.dump(CLEAN, open(pk('clean'), 'wb')); print('준비 완료', len(df), '행', CLEAN, flush=True)
        if A.stage == 'prep': return
    df = pd.read_pickle(pk('prep')); CLEAN.update(pickle.load(open(pk('clean'), 'rb'))); last = int(df['연도'].max()); years = list(range(last - 4, last + 1))
    if A.stage == 'final':
        val_zone, val_cal, val_sweep, val_sum = pickle.load(open(pk('val'), 'rb'))
    else:
        val_zone, val_cal, val_sweep, val_sum = validate(df, years)
        pickle.dump((val_zone, val_cal, val_sweep, val_sum), open(pk('val'), 'wb'))
        if A.stage == 'val': return
    final(A, df, years, last, val_zone, val_cal, val_sweep, val_sum)


def validate(df, years):

    # 검증 두 번
    val_zone, val_cal, val_sweep, val_sum = [], [], [], []
    for k in (2, 1):
        tr = years[:-k]; te = years[-k]
        for adjust in (False, True):
            U, sh, _ = fit_all(df, tr, adjust); T = predict(df, U, sh, te)
            lab = '보정 후' if adjust else '보정 전'
            z = zone_table(T); z.insert(0, '방식', lab); z.insert(0, '검증연도', te); val_zone.append(z)
            val_sum.append(dict(검증연도=te, 방식=lab, 인원=len(T), 예측평균=T.p.mean(), 실제합격률=T.y.mean(), AUC=auc(T.p, T.y), Brier=((T.p - T.y) ** 2).mean()))
            T['bin'] = pd.cut(T.p, np.linspace(0, 1, 11), include_lowest=True)
            c = T.groupby('bin', observed=True).agg(인원=('y', 'size'), 예측평균=('p', 'mean'), 실제합격률=('y', 'mean')).reset_index()
            c['bin'] = c['bin'].astype(str); c.insert(0, '방식', lab); c.insert(0, '검증연도', te); val_cal.append(c)
            if adjust:
                for t in [.1, .15, .2, .3, .4, .5, .6, .7, .75, .8, .85]:
                    s = T[T.p >= t]; val_sweep.append(dict(검증연도=te, 기준=t, 해당인원=len(s), 실제합격률=s.y.mean(), 예측평균=s.p.mean()))
        print('검증', te, '완료', flush=True)
    return val_zone, val_cal, val_sweep, val_sum


def final(A, df, years, last, val_zone, val_cal, val_sweep, val_sum):
    U, shift, D = fit_all(df, years, True); R = unit_rows(df[df['연도'].isin(years)], U)
    U['만점'] = U['과목'].map(SET_MAX)
    for t in [.15, .4, .5, .7]:
        v = ((logit(t) - U.a - shift) / U.b).round(0)
        U[f'{int(t*100)}% 점수'] = v.where((v >= 0) & (v <= U['만점']))   # 만점을 넘으면 비움 = 그 점수로는 도달 못 함
    U['구별력'] = np.where(U.AUC >= .75, '높음', np.where(U.AUC >= .65, '보통', '낮음(점수로 잘 안 갈림)'))
    U.loc[U['50% 점수'].isna(), '구별력'] = '낮음(점수로 잘 안 갈림)'
    bins = {}
    for i in U.index:
        m = R.u.values == i; s = R.s.values[m]; y = R.y.values[m]; yr = R.yr.values[m]
        edges = np.arange(np.floor(s.min() / 20) * 20, np.ceil((s.max() + 1) / 20) * 20 + 1, 20); idx = np.digitize(s, edges) - 1
        bins[i] = [[int(edges[k] + 10), int((idx == k).sum()), int(y[idx == k].sum())] for k in range(len(edges) - 1) if (idx == k).sum()]
        U.loc[i, '최근2년'] = int((yr >= last - 1).sum())
    print('최종 곡선', len(U), '학교', U['학교'].nunique(), '보정값', round(shift, 3))

    write_xlsx(A, U, shift, D, years, pd.concat(val_zone), pd.concat(val_cal), pd.DataFrame(val_sweep), pd.DataFrame(val_sum))
    write_js(A, U, shift, bins, years)


def write_js(A, U, shift, bins, years):
    import json
    data = dict(version=A.version, built=datetime.date.today().isoformat(), years=[years[0], years[-1]], shift=shift, clean=CLEAN,
                zones=[[lo, hi, n] for lo, hi, n in ZONES],
                units=[dict(s=u['학교'], f=u['학부'], k=u['계열'], set=u['과목'], max=int(u['만점']), r=u['회차'], a=round(u.a, 6), b=round(u.b, 6), n=int(u.n), p=int(u['합격']), yk=int(bool(u.get('요강', False))),
                            auc=round(u.AUC, 3), recent=int(u['최근2년']), bins=bins[i]) for i, u in U.iterrows()])
    os.makedirs(os.path.dirname(A.js), exist_ok=True)
    with open(A.js, 'w', encoding='utf-8') as f:
        f.write('/* 자동 생성 — tools/curves.py. 학생 개인 정보 없음(점수 구간별 인원만) */\nwindow.CURVES = ' + json.dumps(data, ensure_ascii=False) + ';\n')


def write_xlsx(A, U, shift, D, years, vz, vc, vs, vsum):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); H = PatternFill('solid', fgColor='27406B'); HF = Font(color='FFFFFF', bold=True)

    def sheet(title, frame, widths=None, pct=(), num1=()):
        ws = wb.create_sheet(title); ws.append(list(frame.columns))
        for c in ws[1]: c.fill = H; c.font = HF; c.alignment = Alignment(horizontal='center', wrap_text=True)
        for row in frame.itertuples(index=False): ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
        for j, col in enumerate(frame.columns, 1):
            L = get_column_letter(j); ws.column_dimensions[L].width = (widths or {}).get(col, max(10, min(28, len(str(col)) * 2 + 2)))
            for cell in ws[L][1:]:
                if col in pct: cell.number_format = '0.0%'
                elif col in num1: cell.number_format = '0.000'
        ws.freeze_panes = 'A2'; return ws

    ws = wb.active; ws.title = '설명'
    lines = [
        f'EJU 합격 곡선 {A.version} — 만든 날 {datetime.date.today().isoformat()}',
        '',
        '■ 무엇인가',
        f'  {years[0]}~{years[-1]}년 지원자 데이터로, 학교·학부·계열마다 「EJU 총점 → 합격률」 S자 곡선(로지스틱 회귀)을 만든 표입니다.',
        f'  곡선 {len(U)}개 · 학교 {U["학교"].nunique()}곳 · 학교·학부 {U[["학교","학부"]].drop_duplicates().shape[0]}개. 합격·불합격이 각각 {MIN_EACH}명 이상인 곳만 만들었습니다.',
        '',
        '■ 데이터 정리 (v2)',
        '  여러 학원이 기관에 제출한 자료라 믿기 어려운 기록이 섞여 있어, 곡선을 만들기 전에 아래를 제거했습니다.',
        f'  · 같은 줄이 통째로 중복된 기록 {CLEAN.get("중복", 0):,}건 제거',
        f'  · 영어 필수 학부({CLEAN.get("영어필수_학부", 0)}곳: 합격자의 {ENG_SHARE:.0%} 이상이 TOEFL·TOEIC 기재)인데 영어 증빙이 없는 기록 {CLEAN.get("영어증빙없음", 0):,}건 제거 (그중 합격 {CLEAN.get("영어증빙없음_합격", 0):,}건)',
        f'  · 학부·합격여부별로 점수가 중앙값에서 {OUT_Z}표준편차 상당 이상 벗어난 점수 {CLEAN.get("점수이상치(점수 칸)", 0):,}칸은 계산에서 제외 (기록 전체는 유지)',
        f'  · 처음 {CLEAN.get("처음", 0):,}건 → 정리 후 {CLEAN.get("남음", 0):,}건',
        '',
        '■ 어떤 점수로 비교하나 (기술 제외)',
        '  계열: 종합과목을 본 지원자 = 문과, 이과 2과목을 본 지원자 = 이과, 일본어만 본 지원자 = 일본어만.',
        '  과목: 학부마다 실제로 보는 과목이 달라서(예: 호세이 경영은 일본어만 봄), 후보 조합(일본어 / +종합 / +수학 등)과',
        '        회차(6월·11월·둘 중 높은 점수) 중에서 그 학부의 합격·불합격을 가장 잘 가르는 것을 자동으로 골랐습니다 → 「과목」「사용 회차」 칸.',
        '  모집요강과 다르면 모집요강이 맞습니다. 이 칸은 「합격 데이터로 본 실제 판정 기준」의 추정입니다.',
        '',
        '■ 연도 보정 (중요)',
        '  같은 점수의 합격률이 해마다 떨어지고 있습니다 (문과+수학 600~650점대: 2022 47% → 2026 36%).',
        '  그래서 최근 연도에 가중치(1·0.7·0.5·0.35·0.25)를 주고, 연도별 공통 보정값 δ를 함께 추정해 가장 최근 해의 수준으로 맞췄습니다.',
        '  연도별 δ (클수록 같은 점수에서 잘 붙던 해): ' + ', '.join(f'{k}: {v:+.2f}' for k, v in D.items()),
        f'  다음 해 예측에 쓰는 보정값 = {shift:+.3f} (가장 최근 해). 추세가 계속되면 실제가 이보다 조금 낮을 수 있습니다.',
        '',
        '■ 검증 (「검증_요약」「검증_구간」「검증_보정」「기준점」 시트)',
        '  미래를 모르는 상태를 만들어 두 번 검증: 2022~2024 → 2025 예측, 2022~2025 → 2026 예측.',
        '  보정 전에는 모든 구간에서 예측이 실제보다 높았고, 보정 후에는 대체로 맞습니다.',
        '  80% 이상 예측 구간에서도 실제 합격률은 75~80% 근처에서 멈춥니다 → 점수만으로는 설명되지 않는 요소(서류·면접 등)가 있다는 뜻.',
        '',
        '■ 읽는 법 (「곡선」 시트)',
        '  50% 점수 = 반반 점수(사실상 합격선). 70%·40%·15% 점수 = 안정·적정·도전 구간이 시작되는 점수.',
        '  칸이 비어 있으면 그 과목 만점으로도 그 확률에 닿지 않는다는 뜻입니다(곡선이 완만 = 점수 외 요소가 큼).',
        '  구별력(AUC) = 점수로 합격·불합격이 얼마나 갈리는지. 「낮음」인 학부는 EJU 외 요소가 크니 곡선을 참고만 하세요.',
        '  「계산기」 시트에서 학교·학부·계열과 점수를 넣으면 합격 가능성이 나옵니다.',
        '',
        '■ 한계',
        '  · 이 데이터는 점수를 알려준 지원자만의 기록입니다(전체 지원자 아님).',
        '  · 학과 차이는 무시하고 학부 단위로 묶었습니다. 같은 학부라도 학과에 따라 합격선이 다를 수 있습니다.',
        '  · 합격 가능성은 「과거 비슷한 점수 지원자 중 합격한 비율」입니다. 합격을 보장하지 않습니다.',
    ]
    for l in lines: ws.append([l])
    ws.column_dimensions['A'].width = 140; ws['A1'].font = Font(bold=True, size=14)

    cols = ['학교', '학부', '계열', '과목', '회차', 'n', '합격', '합격률', '최근2년', 'AUC', '구별력', '15% 점수', '40% 점수', '50% 점수', '70% 점수', 'a', 'b', '키']
    out = U.copy(); out['합격률'] = out['합격'] / out['n']; out['키'] = out['학교'] + '|' + out['학부'] + '|' + out['계열']
    out['과목'] = out['과목'].map(SET_LABEL)
    out = out.sort_values(['학교', '학부', '계열'])[cols].rename(columns={'회차': '사용 회차', 'n': '학습 인원', '최근2년': '최근 2년 인원'})
    sheet('곡선', out, {'학교': 26, '학부': 18, '계열': 11, '구별력': 22, '키': 40, 'a': 12, 'b': 12}, pct=('합격률',), num1=('AUC',))

    wc = wb.create_sheet('계산기')
    rows = [['항목', '입력 / 결과', '설명'], ['학교', '法政大学', '곡선 시트의 학교 이름 그대로 (일본어)'], ['학부', '経営', '곡선 시트의 학부 이름 그대로'],
            ['계열', '문과', '문과 / 이과 / 일본어만'], ['예상 점수', 330, '곡선 시트 「과목」 칸에 적힌 과목의 합 (기술 제외)'],
            ['연도 보정값', round(shift, 4), '다음 해 예측용 (설명 시트 참고)'], ['', '', ''],
            ['합격 가능성', '=IFERROR(1/(1+EXP(-(INDEX(곡선!P:P,MATCH(B2&"|"&B3&"|"&B4,곡선!R:R,0))+INDEX(곡선!Q:Q,MATCH(B2&"|"&B3&"|"&B4,곡선!R:R,0))*B5+B6))),"곡선 없음")', ''],
            ['구간', '=IF(ISNUMBER(B8),IF(B8>=0.7,"안정",IF(B8>=0.4,"적정",IF(B8>=0.15,"도전","15% 미만"))),"")', '초안 기준 70/40/15'],
            ['50% 점수', '=IFERROR(INDEX(곡선!N:N,MATCH(B2&"|"&B3&"|"&B4,곡선!R:R,0)),"")', '이 학부의 반반 점수']]
    for r in rows: wc.append(r)
    for c in wc[1]: c.fill = H; c.font = HF
    wc['B8'].number_format = '0.0%'; wc['B8'].font = Font(bold=True, size=14)
    for c in ['B2', 'B3', 'B4', 'B5']: wc[c].fill = PatternFill('solid', fgColor='FFF2CC')
    wc.column_dimensions['A'].width = 14; wc.column_dimensions['B'].width = 24; wc.column_dimensions['C'].width = 50

    sheet('검증_요약', vsum, pct=('예측평균', '실제합격률'), num1=('AUC', 'Brier'))
    sheet('검증_구간', vz, pct=('예측평균', '실제합격률'))
    sheet('검증_보정', vc, {'bin': 18}, pct=('예측평균', '실제합격률'))
    sheet('기준점', vs.rename(columns={'기준': '예측 ≥ 기준'}), pct=('예측 ≥ 기준', '실제합격률', '예측평균'))
    os.makedirs(os.path.dirname(os.path.abspath(A.xlsx)), exist_ok=True); wb.save(A.xlsx)
    print('엑셀 저장', A.xlsx)


if __name__ == '__main__':
    main()
